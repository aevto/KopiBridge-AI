from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "preliminary-report.md"
OUTPUT = ROOT / "KopiBridge-AI-Preliminary-Project-Report.docx"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
MUTED = RGBColor(89, 89, 89)
BLACK = RGBColor(0, 0, 0)


def set_run_font(run, name="Calibri", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_style_font(style, size, color=BLACK, bold=False, name="Calibri"):
    font = style.font
    font.name = name
    font.size = Pt(size)
    font.color.rgb = color
    font.bold = bold
    style.element.rPr.rFonts.set(qn("w:ascii"), name)
    style.element.rPr.rFonts.set(qn("w:hAnsi"), name)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    set_run_font(run, size=9, color=MUTED)
    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_1)
    run._r.append(instr_text)
    run._r.append(fld_char_2)


def next_numbering_id(numbering, tag_name):
    values = []
    for item in numbering.findall(qn(tag_name)):
        attr_name = "w:abstractNumId" if tag_name.endswith("abstractNum") else "w:numId"
        value = item.get(qn(attr_name))
        if value and value.isdigit():
            values.append(int(value))
    return (max(values) + 1) if values else 1


def create_numbering_definition(doc):
    numbering = doc.part.numbering_part.element
    abstract_id = next_numbering_id(numbering, "w:abstractNum")
    num_id = next_numbering_id(numbering, "w:num")

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))

    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)

    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")

    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    lvl.append(start)

    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "decimal")
    lvl.append(num_fmt)

    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "%1.")
    lvl.append(lvl_text)

    lvl_jc = OxmlElement("w:lvlJc")
    lvl_jc.set(qn("w:val"), "left")
    lvl.append(lvl_jc)

    ppr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "720")
    tabs.append(tab)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "720")
    ind.set(qn("w:hanging"), "360")
    ppr.append(tabs)
    ppr.append(ind)
    lvl.append(ppr)

    abstract.append(lvl)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_num_id = OxmlElement("w:abstractNumId")
    abstract_num_id.set(qn("w:val"), str(abstract_id))
    num.append(abstract_num_id)
    numbering.append(num)

    return num_id


def apply_numbering(paragraph, num_id):
    ppr = paragraph._p.get_or_add_pPr()
    num_pr = ppr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        ppr.append(num_pr)

    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_el = OxmlElement("w:numId")
    num_id_el.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num_id_el)


def apply_doc_styles(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles

    normal = styles["Normal"]
    set_style_font(normal, 11, BLACK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    h1 = styles["Heading 1"]
    set_style_font(h1, 16, BLUE, bold=True)
    h1.paragraph_format.space_before = Pt(16)
    h1.paragraph_format.space_after = Pt(8)
    h1.paragraph_format.keep_with_next = True

    h2 = styles["Heading 2"]
    set_style_font(h2, 13, BLUE, bold=True)
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(6)
    h2.paragraph_format.keep_with_next = True

    h3 = styles["Heading 3"]
    set_style_font(h3, 12, DARK_BLUE, bold=True)
    h3.paragraph_format.space_before = Pt(8)
    h3.paragraph_format.space_after = Pt(4)
    h3.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Number"):
        style = styles[style_name]
        set_style_font(style, 11, BLACK)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.167
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)

    if "Reference" not in styles:
        styles.add_style("Reference", 1)
    reference = styles["Reference"]
    set_style_font(reference, 10, BLACK)
    reference.paragraph_format.space_before = Pt(0)
    reference.paragraph_format.space_after = Pt(7)
    reference.paragraph_format.line_spacing = 1.08
    reference.paragraph_format.first_line_indent = Inches(-0.35)
    reference.paragraph_format.left_indent = Inches(0.35)


def add_header_footer(section):
    header = section.header.paragraphs[0]
    header.text = ""
    run = header.add_run("KopiBridge AI Preliminary Project Report")
    set_run_font(run, size=9, color=MUTED)
    header.paragraph_format.space_after = Pt(0)

    footer = section.footer.paragraphs[0]
    footer.text = ""
    add_page_number(footer)


def clean_inline(text):
    text = text.replace("O*NET", "O__STAR__NET")
    text = text.replace("O\\*NET", "O__STAR__NET")
    text = text.replace("`", "")
    text = re.sub(r"\*\*([^*]+)\*\*", r"\1", text)
    text = re.sub(r"\*([^*]+)\*", r"\1", text)
    return text.replace("O__STAR__NET", "O*NET").strip()


def add_paragraph_with_soft_breaks(doc, text, style=None):
    p = doc.add_paragraph(style=style)
    for index, part in enumerate(text.split("  ")):
        if index:
            p.add_run().add_break()
        run = p.add_run(clean_inline(part))
        set_run_font(run, size=11 if style != "Reference" else 10, color=BLACK)
    return p


def add_numbered_paragraph(doc, text, num_id):
    p = doc.add_paragraph()
    apply_numbering(p, num_id)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.167
    run = p.add_run(clean_inline(text))
    set_run_font(run, size=11, color=BLACK)
    return p


def add_cover(doc):
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(74)

    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = kicker.add_run("PRELIMINARY PROJECT REPORT")
    set_run_font(run, size=10, color=BLUE, bold=True)
    kicker.paragraph_format.space_after = Pt(16)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("KopiBridge AI")
    set_run_font(title_run, size=28, color=BLACK, bold=True)
    title.paragraph_format.space_after = Pt(6)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.add_run("Resume-to-AI-Tech-Role Gap Analyser")
    set_run_font(subtitle_run, size=14, color=MUTED)
    subtitle.paragraph_format.space_after = Pt(30)

    metadata = [
        ("Student name", "[Your Name]"),
        ("Module", "Final Year Project"),
        ("Project template", "Software Development / Web Application template"),
        ("Date", "26 June 2026"),
        ("Approximate word count", "4,300 words excluding references"),
    ]

    for label, value in metadata:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        label_run = p.add_run(f"{label}: ")
        set_run_font(label_run, size=11, color=BLACK, bold=True)
        value_run = p.add_run(value)
        set_run_font(value_run, size=11, color=BLACK)
        p.paragraph_format.space_after = Pt(4)

    doc.add_page_break()


def build_docx():
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = Document()
    apply_doc_styles(doc)
    add_header_footer(doc.sections[0])
    add_cover(doc)

    in_body = False
    in_references = False
    first_chapter = True
    current_num_id = None

    for raw in lines:
        line = raw.rstrip()
        if line.startswith("## Chapter 1:"):
            in_body = True

        if not in_body:
            continue

        if not line.strip():
            current_num_id = None
            continue

        if line.startswith("## "):
            current_num_id = None
            title = clean_inline(line[3:])
            if title == "References":
                in_references = True
                doc.add_page_break()
                doc.add_heading(title, level=1)
                continue
            if not first_chapter:
                doc.add_page_break()
            first_chapter = False
            doc.add_heading(title, level=1)
            continue

        if line.startswith("### "):
            current_num_id = None
            doc.add_heading(clean_inline(line[4:]), level=2)
            continue

        numbered = re.match(r"^\d+\.\s+(.*)$", line)
        if numbered:
            if current_num_id is None:
                current_num_id = create_numbering_definition(doc)
            add_numbered_paragraph(doc, numbered.group(1), current_num_id)
            continue

        current_num_id = None

        if in_references:
            add_paragraph_with_soft_breaks(doc, line, style="Reference")
            continue

        add_paragraph_with_soft_breaks(doc, line)

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_docx()
