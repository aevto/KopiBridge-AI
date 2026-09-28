"""Revise the retained report using measured evaluation records and real screenshots."""
from pathlib import Path
import json, re, shutil, statistics, textwrap
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[1]
REPORTS = ROOT / 'reports'
OUT = REPORTS / 'KopiBridge-AI-Final-Project-Report.docx'
BACKUP = REPORTS / 'KopiBridge-AI-Final-Project-Report-Before-Multimodal.docx'
ASSETS = REPORTS / 'final-report-assets/multimodal'
if not BACKUP.exists(): shutil.copy2(OUT, BACKUP)
d = Document(BACKUP)
old = list(d.paragraphs)
runs = json.loads((REPORTS/'evaluation/model-results.json').read_text())['runs']
browser = json.loads((REPORTS/'evaluation/browser-results.json').read_text())
assert browser['checks']['deletion'] and not browser['checks']['consoleErrors']

replace = {
1: 'KopiBridge AI',
2: 'Multimodal Evidence Led Career Guidance',
3: 'Final Project Report',
4: 'CM3020 Artificial Intelligence\nProject Idea 1 Section 4.1\nOrchestrating AI models to achieve a goal',
5: 'Code repository https://github.com/aevto/KopiBridge-AI\nExisting deployment https://kopi-bridge-ai.vercel.app',
6: '28 September 2026',
9: 'KopiBridge AI helps early-career applicants compare resume evidence with a target role and prepare truthful next steps. Its final implementation orchestrates three distinct pretrained models: image-to-text extraction, text guidance and speech transcription. Scanned resume pages become editable text; a deterministic engine builds the authoritative evidence report; a guarded language model refines permitted guidance; and a transcribed interview answer receives feedback tied to that same report. The applicant reviews both extracted text and transcription before relying on them. The product includes authenticated accounts, private history, server-enforced analysis credits and print export. Evaluation combines executable database tests, browser integration checks and exploratory comparisons of six model configurations across synthetic inputs. Initial feedback quotation failures motivated a schema-constrained citation design, which passed all six subsequent quotation checks. These results establish integration feasibility, not hiring prediction accuracy or proven career benefit. Important limitations include heuristic negation errors, a small synthetic evaluation set, compressed evidence context and the absence of a participant study. The contribution is a coherent multimodal preparation workflow that keeps claim status inspectable rather than delegating it to generated prose.',
18: 'The project began with local-only analysis and subsequently added authenticated history, transactional credits and guarded text refinement. The final extension addresses the confirmed AI orchestration template through image, text and audio models working towards one career-preparation goal. Without an OpenAI key, pasted text and selectable PDFs still produce the complete deterministic report. Image reading and transcription require provider access; unavailable interview feedback falls back to an explicitly labelled checklist. These fallbacks preserve utility but are not represented as a successful three-model run.',
28: 'Evaluate the integrated software and each model with reproducible fixtures, component-appropriate measures and explicit limitations.',
29: 'Three questions guide evaluation: (RQ1) Can distinct image, text and audio models form a coherent resume-to-interview workflow? (RQ2) Can deterministic evidence labels and grounded quotations prevent unsupported changes while allowing useful guidance? (RQ3) Do authentication, ownership, credits and recovery controls work across the integrated path?',
31: 'The confirmed template is CM3020 Artificial Intelligence, Project Idea 1, section 4.1: Orchestrating AI models to achieve a goal. Its final-product requirement is at least three pretrained models operating on different data types or domains, combined in a purposeful working system. KopiBridge AI uses gpt-4.1-mini for resume images, gpt-5.6-terra for text guidance and gpt-4o-mini-transcribe for interview audio. PDF.js and the deterministic matching engine are supporting software, not additional pretrained models. The same saved report connects the three model roles, rather than presenting independent demonstrations.',
32: 'The final scope includes selectable and scanned PDFs, image uploads, editable extracted text, local evidence analysis, guarded narrative guidance, short audio answers, reviewed transcripts and private interview feedback. Employer-side selection, speech-based personality assessment, automated applications, public reports and LLM-generated scores are excluded. The existing text product is deployed; the multimodal extension evaluated here runs locally against live providers and the approved hosted database migration. It has not been claimed as a newly deployed production release.',
57: 'Resume documents and interview answers can contain personal information. Selectable PDF text is extracted locally; scanned pages and audio require explicit consent before provider processing. KopiBridge does not persist those raw uploads. Reports retain short evidence excerpts and reviewed interview answers, so they remain personal data even though complete original documents are not stored. Deleting a report cascades to its interview feedback, while minimal usage events remain for quota enforcement. Provider retention is separate: disabling retrievable response storage is not a guarantee of zero provider retention.',
63: 'KopiBridge AI combines applicant-side evidence analysis with a multimodal preparation workflow. Its contribution is integration and constraint design, not training a new foundation model. The relevant test is whether pretrained capabilities improve access and preparation while keeping evidence, ownership and uncertainty visible. The following implementation and evaluation therefore distinguish software invariants, extraction accuracy, quotation grounding and user usefulness.',
68: 'A representative user uploads a selectable PDF or scanned resume, reviews the extracted text, pastes a vacancy and confirms its role metadata. After generating a private report, the user reviews the most important gaps and chooses one interview question. A short recorded or uploaded answer is transcribed, corrected and submitted for feedback grounded in that report. The applicant can revisit the combined result or print it, and later repeat analysis after producing genuine new evidence.',
82: 'Raw PDFs, images and audio are not persisted by KopiBridge. History contains report JSON and role metadata. Interview practice adds the selected question, reviewed transcript, structured feedback, source, model and timestamp. Audit records hold operation identifiers and Singapore usage dates, not raw media. These are deliberate retention boundaries, not a claim that the stored report is anonymous.',
88: 'Five tables support the final workflow: analyses, daily_credit_balances, credit_transactions, model_operations and interview_practice. The original report and credit structures are preserved. New model-operation records enforce six daily attempts separately for vision, transcription and interview feedback; practice records reference both an owned report and accepted operation. Row-level security restricts reads and deletions to the owner. Raw upload storage is absent.',
95: 'The workplan progressed from the local engine to product infrastructure, guarded text refinement and multimodal integration. The approved additive migration and local integration tests complete the implemented scope. Deployment of the new extension, representative human evaluation and stronger provider-independent validation remain distinct next steps. No unrecorded tutor feedback or participant study is assumed; the revised scope follows the confirmed section 4.1 template and observed test failures.',
105: 'ResumeUpload accepts PDF, PNG, JPEG and WebP with an 8MB client limit. PDF.js first attempts selectable-text extraction for up to twenty pages. A low-text PDF offers consented image reading for at most three pages. Pages are rasterised in-browser and bounded before transmission. The server validates format signatures, payload length and consent, then calls the vision model. Extracted text remains editable and warnings are shown. An upload control is disabled until hydration so an early selection cannot be lost before its event handler is ready.',
106: 'This retains the low-cost selectable-PDF path while adding a meaningful vision requirement. Image requests have a 3.8MB server-body limit and are read incrementally, including when Content-Length is absent. Raw file data is transient. A sample and pasted text remain alternatives when extraction fails. Selecting a new valid file clears the prior text to avoid accidentally analysing a different resume; a failed analysis request, by contrast, preserves the current input.',
175: 'The project remains one typed Next.js codebase. The additive migration 20260927203904_add_multimodal_career_workflow.sql extends, rather than rewrites, the foundation. Model IDs are configurable server-side and checked for distinctness. Provider requests have bounded timeouts and no automatic retry multiplication. Reproducible synthetic fixtures, model-result JSON and a disposable-account browser runner support regression checks; the production app itself does not require an administrator key.',
179: 'Evaluation addresses model orchestration as well as the web product. Deterministic and schema tests cover invariant preservation; executable SQL tests cover ownership and quota transitions; live model comparisons measure extraction and grounding; and browser tests connect the modalities through a saved report. This combination addresses RQ1-RQ3 without treating a polished screen or one percentage as evidence of career effectiveness.',
182: 'Testing follows a risk-based plan (Table 3). Each case states its input, observable outcome and evidence level. Invalid requests should be rejected before processing, models must not change authoritative fields, and new media must not spend a full-analysis credit. Browser checks exercise the actual local application with live Supabase and OpenAI, using disposable synthetic accounts. Database tests run separately in PGlite; they are not labelled hosted stress tests.',
186: 'The current Vitest run passed 33 tests in eight files. Existing tests cover analysis, credit dates, validation and the narrative merge. New tests execute both SQL migrations in PGlite and check media quotas, duplicate keys, daily-date partitioning, owner isolation, allowed questions, save idempotency, deletion cascade and credit refund behaviour. HTTP tests verify authenticated identity, same-host origin handling and streamed size limits. Media tests reject malformed images, overlong or silent WAV data, injected identity fields and invented quotations. The constrained feedback schema accepts only transcript excerpts.',
187: 'Six public Playwright checks passed across desktop and mobile profiles. A separate authenticated integration case passed against the local application and live services, using two disposable accounts removed afterwards. Table 4 summarises the final command checks. The earlier chart in Figure 24 is retained as historical text-product evidence and must not be read as the current test count.',
192: 'The final authenticated case uploaded a selectable PDF, then an image-only PDF, obtained consented vision extraction and created a report. It uploaded a synthetic spoken answer, received a transcription, confirmed its text and saved model feedback. Reloading the report recovered the feedback. The daily balance fell from three to two after analysis and stayed at two after transcription and feedback. A second user received an unavailable-report page and HTTP 404 for foreign feedback and deletion requests.',
193: 'The same run submitted concurrent full-analysis requests with one idempotency key: exactly one returned HTTP 201 and the other returned HTTP 409. A third successful analysis exhausted the allowance; the fourth returned HTTP 429 and the UI disabled analysis. Owner deletion required confirmation and succeeded. Window.print invocation was observed through a test replacement, and Chromium produced a PDF for visual inspection. Screenshots are from the actual local application, not design mock-ups.',
196: 'PGlite provides executable PostgreSQL semantics for the migrations but uses a single local connection, so its overlapping promises are not evidence of distributed concurrency. The hosted duplicate-request case gives narrower integration evidence, not a load benchmark. Supabase advisors reported intentional authenticated security-definer RPCs, disabled leaked-password protection and limited MFA options. Ownership checks and fixed search paths were reviewed, but these warnings and direct owner-callable RPCs mean the product should not be described as independently security-certified or its reports as tamper-proof credentials.',
198: 'Earlier figures show the established 70% sample report. The new scanned fixture deliberately includes a negated statement: no Docker, cloud deployment or RAG project yet. The unchanged lexical matcher nevertheless assigned partial evidence to some of those terms. This is a substantive failure, not an OCR error: the image model reproduced the words correctly. Preserving deterministic fields limits model authority but also preserves deterministic mistakes. The report must therefore remain an inspectable aid; a future negation-aware evidence benchmark is higher priority than claiming calibrated accuracy.',
199: 'The guarded text layer preserved all tested scores, evidence-map entries and gap severities for both candidate models. That is a contract result, not a proof of truthful prose. Inspection found that the smaller text candidate suggested inserting job-description terms too broadly, while feedback sometimes treated details absent from the compressed evidence context as unverified even when present in the full resume. These observations motivate explicit claim-level validation and better evidence selection. No precision, recall or employment-outcome statistic is inferred from the small fixtures.',
200: 'The browser case found no horizontal overflow at widths 390, 768 and 1440 pixels and no captured console or page errors. The report remained readable with distinct claim labels. Print controls were hidden and a PDF was generated; Safari, unusual paper sizes and very long reports remain outside the observed scope. Microphone recording is implemented, but this automated end-to-end case tested the audio-upload path rather than a physical microphone.',
209: 'The strongest result is an integrated three-model path with preserved local evidence, private persistence and recoverable model failures. Quotation validation initially rejected five of six model responses; constraining allowed excerpts then produced six accepted reruns without loosening the verification rule. This is concrete iteration based on failure evidence. It does not guarantee that every interpretation of a correctly quoted sentence is sound.',
210: 'The main weaknesses are lexical negation errors, compressed evidence context and evaluation breadth. Synthetic speech has limited accent and delivery diversity, and three resume layouts cannot represent real document noise. No participant usability scores or recruitment outcomes were collected. Hosted configuration still needs stronger account protections, and a future design should prevent clients from directly supplying apparently model-produced report content through owner-callable write RPCs.',
211: 'Next priorities are a human-labelled evidence set with negation cases, a consented speech/document corpus, repeated candidate-model runs and the formative study described above. Operational work includes recording-permission tests, multi-browser print checks, least-privilege server-only finalisation and account-wide deletion. These improvements follow observed limitations rather than adding unrelated features.',
214: 'KopiBridge AI implements CM3020 Project Idea 1, section 4.1, through three distinct pretrained models handling images, text and audio within one career-preparation workflow. Scanned documents become reviewed resume text; a deterministic report anchors language guidance; and a reviewed interview transcript receives private feedback tied to that report. The original account, credit, history and print functions remain part of the integrated product.',
215: 'The contribution is the orchestration and evidence policy rather than a new foundation model. Explicit review points, constrained response schemas and immutable local fields reduce the opportunity for generated prose to become unsupported authority. The quotation redesign shows how evaluation can lead to a concrete implementation improvement. At the same time, a preserved heuristic can still be wrong, so inspectability must not be confused with semantic correctness.',
216: 'The final evidence includes 33 passing automated tests, six passing public browser checks, one passing authenticated end-to-end workflow and 26 recorded live model trials. The selected vision model had zero normalised word error on three synthetic layouts; the selected speech model had zero error on two answers and one word error on the third. Both text candidates preserved fixed assessment fields. Five initial interview outputs were rejected for quotation mismatch, followed by six valid constrained reruns. These small results support technical feasibility, not general model superiority or a measured improvement in employability.',
217: 'The multimodal code is validated locally against live services and its approved database migration is applied. Publishing that extension remains separate from implementation. The public repository must contain the submitted revision and remain accessible until results are issued. Further research should test real users, diverse speech and document layouts, claim-level correctness and the effect of completed proof tasks. No unperformed user study is presented as a finding.',
}
for index, text in replace.items(): old[index].text = text
old[1].style = d.styles['Title']
old[1].runs[0].font.color.rgb = RGBColor(0,0,0)
for i in [2,3,4,5,6]:
    for r in old[i].runs: r.font.size = Pt(12 if i > 2 else 16)

def before(anchor, text='', style='Normal'):
    return anchor.insert_paragraph_before(text, style)
def heading(anchor, text): return before(anchor, text, 'Heading 2')
def picture(anchor, key, filename, caption, width=6.3):
    p = before(anchor)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(ASSETS/filename), width=Inches(width))
    # Prevent portrait screenshots from exceeding printable height.
    shape = d.inline_shapes[-1]
    if shape.height > Inches(6.5):
        ratio = Inches(6.5)/shape.height
        shape.width = int(shape.width*ratio); shape.height = Inches(6.5)
    before(anchor, f'Figure {key}. {caption}', 'Caption')
def table(anchor, title, rows):
    before(anchor, title, 'Caption').paragraph_format.keep_with_next = True
    t = d.add_table(rows=1, cols=len(rows[0])); t.style = 'Table Grid'
    anchor._p.addprevious(t._tbl)
    for c, val in zip(t.rows[0].cells, rows[0]): c.text = str(val)
    for row in rows[1:]:
        for c, val in zip(t.add_row().cells, row): c.text = str(val)
    for c in t.rows[0].cells:
        for r in c.paragraphs[0].runs:r.bold=True
    before(anchor).paragraph_format.space_after=Pt(0)
    return t

# Extend the reviewed literature rather than replace it with a list of APIs.
heading(old[61], '2.7 Multimodal extraction and human review')
before(old[61], 'Visual document understanding differs from parsing an existing text layer. Kim et al. (2022) introduce Donut as an OCR-free document-understanding transformer and identify costs and error propagation in OCR-dependent pipelines. This motivates evaluating page images as model inputs, but Donut was not implemented or benchmarked here. A hosted vision model was chosen to minimise separate infrastructure and use the existing provider boundary. The trade-off is external processing, variable cost and potential hallucinated text. A conventional local OCR engine remains a useful future baseline, not a tested rejected alternative.')
before(old[61], 'Radford et al. (2023) demonstrate broad speech-recognition capability from large-scale weak supervision. Their benchmark results justify considering pretrained transcription, not assuming it will accurately capture every applicant or technical term. KopiBridge compares two hosted transcription configurations using word error rate and latency, then requires user correction. Interview content is assessed as an unverified statement; neither accent nor vocal style is used to infer personality or employability. Synthetic speech is suitable for regression fixtures but insufficient for claims about accent fairness or real interview performance.')
before(old[61], 'Amershi et al. (2019) emphasise that human-AI interaction must accommodate imperfect predictions and support correction. The editable extraction and transcription stages apply that principle directly. OpenAI structured outputs constrain response shape (OpenAI, n.d.-c), but shape compliance alone cannot prove factual grounding. KopiBridge therefore combines schema validation with source-excerpt constraints and preserves deterministic claim labels. The model comparison must evaluate these application-specific properties, not only fluency. This synthesis supports separate measures for transcription fidelity, evidence invariants and user comprehension.')
old[61].text = '2.8 Synthesis and project position'

heading(old[100], '3.9 Integrated model orchestration')
before(old[100], 'Figure NEW1 shows the final information flow. The vision model supplies editable text rather than an employment judgement. The text model refines a precomputed report, and the audio model supplies a transcript rather than a score. The reviewed transcript and selected question then return to the text model with the same report evidence. This dependency makes the modalities serve one goal: turning an existing application into a defensible account of readiness.')
before(old[100], 'Each media endpoint authenticates independently, checks ownership where a report is involved, validates payload bounds and reserves an operation. An advisory transaction lock serialises each user and operation kind; a unique user-kind-request key prevents repeated charging of the same operation. Six attempts per kind per Singapore day are distinct from three full analyses. Failed media attempts count towards the limit to prevent unlimited provider retries. This policy is disclosed before submission and does not refund or alter full-analysis credits.')

# Deterministic raster diagram, drawn from implemented flow rather than a mock UI.
im=Image.new('RGB',(1800,1040),'white'); draw=ImageDraw.Draw(im)
font=ImageFont.truetype('/System/Library/Fonts/Supplemental/Arial.ttf',31)
bold=ImageFont.truetype('/System/Library/Fonts/Supplemental/Arial Bold.ttf',34)
def box(x,y,w,h,title,body,color):
    draw.rounded_rectangle((x,y,x+w,y+h),radius=10,fill=color,outline='#a9b5ae',width=2)
    draw.text((x+25,y+22),title,font=bold,fill='#172a24')
    for j,line in enumerate(textwrap.wrap(body,42)):
        draw.text((x+25,y+78+j*39),line,font=font,fill='#263a33')
def arrow(a,b):
    draw.line((a,b),fill='#41695a',width=5)
    x,y=b; draw.polygon([(x,y),(x-11,y-18),(x+11,y-18)],fill='#41695a')
box(60,35,790,200,'IMAGE   gpt-4.1-mini','Scanned resume pages -> extracted text. Applicant reviews and corrects.', '#f2f6f3')
box(950,35,790,200,'AUDIO   gpt-4o-mini-transcribe','Recorded answer -> transcript. Applicant reviews and confirms.', '#f5f1eb')
box(60,340,790,220,'LOCAL ENGINE   not a pretrained model','Resume + vacancy -> fixed scores, evidence strengths and honesty labels.', '#f7f8fa')
box(950,340,790,220,'TEXT   gpt-5.6-terra','Selected question + reviewed transcript + saved evidence -> grounded feedback.', '#f2f6f3')
box(60,660,790,230,'TEXT   gpt-5.6-terra','Local report -> constrained narrative refinement. Fixed fields cannot change.', '#f2f6f3')
box(950,660,790,230,'PRIVATE REPORT AND PRACTICE','Owner-only persistence, history and print. No raw image or audio storage.', '#f7f8fa')
arrow((450,235),(450,340)); arrow((450,560),(450,660)); arrow((1345,235),(1345,340)); arrow((1345,560),(1345,660))
draw.line(((850,775),(950,775)), fill='#41695a',width=5)
draw.line(((850,450),(950,450)), fill='#41695a',width=5)
draw.text((60,956),'Server boundary: session checks | bounded inputs | consent | usage reservations | no client secrets',font=font,fill='#263a33')
im.save(ASSETS/'model-pipeline.png')
picture(old[100],'NEW1','model-pipeline.png','Final three-model workflow and its deterministic evidence boundary.')

heading(old[177], '4.10 Scanned resumes and spoken interview practice')
before(old[177], 'lib/openai-media.ts supplies server-only adapters for all three media operations. Image extraction uses the Responses API with a bounded text-and-warnings schema. lib/audio.ts decodes accepted audio locally and resamples it to mono 16kHz PCM WAV. The server checks canonical headers, byte counts, duration and non-silent energy before transcription. Audio is limited to ninety seconds; decoding and upload errors leave typed answers available. A selected report question is derived server-side, not accepted as arbitrary client authority.')
before(old[177], 'InterviewPractice provides recording, upload, playback, consent, transcript correction and save states on the existing report page. The model receives the answer and report evidence, not permission to change scores. groundedFeedbackSchema constructs an enumeration of actual transcript spans for quotation fields; validateFeedbackQuotes then checks the returned quotations against the source. The rest of the feedback remains generated interpretation and still requires review. A timeout or invalid response stores a clearly labelled local preparation checklist instead of pretending that AI feedback succeeded.')
before(old[177], 'The approved additive migration creates owner-only interview_practice and model_operations tables. save_interview_practice checks the reservation owner, parent report and selected question before saving. Report deletion cascades to feedback, while quota records remain until account deletion. The implementation uses media request schemas and streamed body limits, and blocks cross-site browser origins while accommodating Next.js localhost host normalisation. Hydration and audio object-URL lifecycle defects found during browser checks were fixed before the successful rerun.')
before(old[177], 'Figures NEW2-NEW7 document the final local workflow using synthetic content and live providers. They show consent, reviewed extraction, the resulting report, audio preparation, corrected transcription and saved feedback. The earlier Figures 5-22 document the established text-product interface; they are retained for coverage, not presented as evidence of the new media routes. The final mobile and credit-limit states are shown in Figures NEW8-NEW9.')
for key,file,caption in [
('NEW2','03-scan-consent.png','Scanned PDF fallback requires explicit consent before provider processing.'),
('NEW3','04-reviewed-resume.png','Vision output is editable and labelled for review before analysis.'),
('NEW4','05-report-summary.png','Actual report created from the scanned fixture, with one analysis credit spent.'),
('NEW5','07-audio-ready.png','Short interview audio can be previewed before transcription.'),
('NEW6','08-transcript-review.png','The applicant reviews the transcript before requesting feedback.'),
('NEW7','09-saved-feedback.png','Saved feedback quotes source excerpts and distinguishes unsupported claims.'),
('NEW8','10-practice-390.png','Final interview workflow at a 390-pixel viewport; no horizontal overflow.'),
('NEW9','12-zero-credits.png','Daily analysis exhaustion produces a clear disabled state.')]: picture(old[177],key,file,caption,3.0 if key=='NEW8' else 6.3)

heading(old[213], '5.9 Model comparison and observed iteration')
before(old[213], 'The reproducible dataset contains three synthetic resume layouts (clean, two-column and low-contrast) and three short synthetic answers (specific, unsupported-claim and noisy). macOS speech synthesis supplies a known reference; fixed-seed noise adds one controlled perturbation. Every candidate sees the same inputs. Word error rate is (substitutions + deletions + insertions) divided by reference words after lowercasing and removing punctuation. This measures content order as well as recognition, but misses punctuation and semantic harm. Latencies are single observations, not service-level estimates.')
rows=[['Task','Candidate','Cases','Mean WER','Mean latency']]
for kind, models in [('vision',['gpt-4.1-mini','gpt-4o-mini']),('audio',['gpt-4o-mini-transcribe','whisper-1'])]:
    for model in models:
        a=[r for r in runs if r['kind']==kind and r['requestedModel']==model and r['success']]
        rows.append([kind,model,len(a),f"{statistics.mean(r['wordErrorRate'] for r in a)*100:.2f}%",f"{statistics.mean(r['latencyMs'] for r in a)/1000:.2f}s"])
table(old[213],'Table 5. Exploratory live extraction comparison on synthetic fixtures.',rows)
before(old[213], 'The selected vision model reproduced all three normalised references. gpt-4o-mini had 18.67% WER on the two-column case and zero on the others, so gpt-4.1-mini was retained for this workflow despite slower average latency. Both transcription candidates tied at 0%, 2.04% and 0% WER; the selected candidate was slightly faster on average. This is weak selection evidence from one synthetic voice, not proof that Whisper is generally inferior. Both text candidates preserved fixed fields, but the smaller model was not promoted because that invariant alone did not establish safer advice.')
before(old[213], 'There were 26 recorded provider trials: six vision, six transcription, two report-guidance, six initial interview-feedback and six revised-feedback calls. Five initial interview responses failed exact-quotation validation. The application rejected them. Constraining quotations to source-span enums yielded six accepted reruns. This demonstrates a useful engineering correction, not zero hallucination: quoted evidence can still be misinterpreted and unseen inputs can fail. Token usage is retained where returned, but no unmeasured monetary cost or statistical significance is claimed.')

# Replace stale test matrices, retaining the established table formatting.
testrows=[
['TP01','Public browser UI','Desktop/mobile entry and validation','6 browser checks','Pass'],
['TP02','Authentication','Login, persisted session, protected routes','Disposable accounts; signup email not rerun','Partial'],
['TP03','Selectable PDF','Local editable extraction','Flask text recovered','Pass'],
['TP04','Scanned PDF','Consent and real vision output','Live extraction to report','Pass'],
['TP05','Audio upload','Verified WAV and transcription','Live speech model','Pass'],
['TP06','Interview feedback','Reviewed answer, source quotes, persistence','Live text model and reload','Pass'],
['TP07','Credit and duplicate','One deduction for concurrent same-key requests','201 and 409; fourth request 429','Pass'],
['TP08','Failure refund','One restoration for pending reservation','Executable SQL; no hosted provider fault injection','Partial'],
['TP09','Day boundary','Singapore dates and independent daily limits','Unit and SQL fixtures','Pass'],
['TP10','Owner isolation','Foreign report, feedback and delete denied','Second account 404','Pass'],
['TP11','Deletion','Confirmation and owner-only deletion','Browser plus SQL cascade test','Pass'],
['TP12','Grounding','Fixed fields preserved; unsupported quotes rejected','Contract tests and 26 model trials','Pass'],
['TP13','Print','window.print called; printable report rendered','Chromium PDF; other browsers untested','Partial'],
['TP14','Responsive','No overflow at 390, 768 and 1440 pixels','Authenticated browser checks','Pass'],
['TP15','Usability benefit','Users understand claims and actionable next steps','Study proposed, not performed','Not tested']]
for row,vals in zip(d.tables[3].rows[1:],testrows):
    for c,val in zip(row.cells,vals): c.text=val
    props=row.cells[-1]._tc.get_or_add_tcPr()
    for shade in list(props.findall(qn('w:shd'))):props.remove(shade)
    shade=OxmlElement('w:shd')
    shade.set(qn('w:fill'),{'Pass':'E8F2EB','Partial':'FFF3DB','Not tested':'F1F1F1'}[vals[-1]])
    props.append(shade)
validation=[['Prettier','Pass','Repository formatting'],['Lint','Pass','ESLint and source security policy'],['TypeScript','Pass','Static checking'],['Vitest','33 tests; 8 files','Local logic, HTTP and executable SQL'],['Playwright','6 public + 1 integrated','Live providers and disposable accounts'],['Production build','Pass','Next.js 16 production compilation'],['npm audit','0 vulnerabilities','Installed dependencies']]
for row,vals in zip(d.tables[4].rows[1:],validation):
    for c,val in zip(row.cells,vals):c.text=val
d.tables[1].rows[2].cells[1].text='PDF/image, pasted text, role metadata and reviewed audio answer'
for vals in [['Multimodal integration','Vision extraction, transcription, grounded feedback','Implemented and locally tested; migration applied'],['Submission','Final report, screenshots, comparison and limitations','Measured evidence; deployment still separate']]:
    for c,val in zip(d.tables[2].add_row().cells,vals):c.text=val

# Keep earlier screenshots identifiable as historical evidence.
for p in d.paragraphs:
    if p.style.name=='Caption' and re.match(r'Figure (?:[5-9]|1[0-9]|2[0-5])\.',p.text):
        p.text += ' Earlier text-product build.'
    if 'single migration for the current foundation' in p.text:
        p.text=p.text.replace('single migration for the current foundation','versioned foundation and additive migrations')

# Make retained hyperlink references readable in exported PDF as well as Word.
for p in old[221:]:
    if 'online source' in p.text:
        links=p._p.xpath('.//w:hyperlink')
        urls=[d.part.rels[e.get(qn('r:id'))].target_ref for e in links if e.get(qn('r:id')) in d.part.rels]
        if urls:p.text=p.text.replace('online source',urls[0])
refs=[
'Amershi, S. et al. (2019) Guidelines for Human-AI Interaction. Proceedings of CHI 2019. doi:10.1145/3290605.3300233. https://www.microsoft.com/en-us/research/publication/guidelines-for-human-ai-interaction/',
'Kim, G. et al. (2022) OCR-free Document Understanding Transformer. Proceedings of ECCV 2022. https://arxiv.org/abs/2111.15664',
'Radford, A., Kim, J.W., Xu, T., Brockman, G., McLeavey, C. and Sutskever, I. (2023) Robust Speech Recognition via Large-Scale Weak Supervision. Proceedings of ICML, PMLR 202, pp. 28492-28518. https://proceedings.mlr.press/v202/radford23a.html',
'OpenAI (n.d.-a) Images and vision. https://developers.openai.com/api/docs/guides/images-vision (Accessed 28 September 2026).',
'OpenAI (n.d.-b) File transcription. https://developers.openai.com/api/docs/guides/speech-to-text (Accessed 28 September 2026).',
'OpenAI (n.d.-c) Structured model outputs. https://developers.openai.com/api/docs/guides/structured-outputs (Accessed 28 September 2026).',
'University of London (n.d.) Final Project Templates. CM3020 Artificial Intelligence, Project Idea 1, section 4.1, pp. 19-21. Course project brief supplied with the project.',
]
for text in refs:d.add_paragraph(text)
reference_paragraphs=[]
in_references=False
for p in d.paragraphs:
    if p.text=='References':in_references=True;continue
    if in_references and p.text.strip():reference_paragraphs.append(p)
sorted_references=sorted((p.text for p in reference_paragraphs),key=str.casefold)
for p,text in zip(reference_paragraphs,sorted_references):p.text=text
before(old[177], 'The provider interfaces follow OpenAI image and file-transcription documentation (OpenAI, n.d.-a; n.d.-b). API availability was checked through actual calls; provider model names identify pretrained services, not locally trained checkpoints.')

# Renumber every figure and its references in document order.
mapping={};n=0
for p in d.paragraphs:
    m=re.match(r'Figure ([A-Z0-9]+)\.',p.text)
    if p.style.name=='Caption' and m:n+=1;mapping[m.group(1)]=str(n)
pattern=re.compile(r'\b(Figures? )([A-Z0-9]+)(?:(-| and )([A-Z0-9]+))?')
def renumber(m):return m.group(1)+mapping.get(m.group(2),m.group(2))+((m.group(3)+mapping.get(m.group(4),m.group(4))) if m.group(4) else '')
for p in d.paragraphs:
    if pattern.search(p.text):p.text=pattern.sub(renumber,p.text)
    if 'Schema changes are committed' in p.text:p.text=p.text.replace('Schema changes are committed','Schema changes are recorded')
    if p.style.name=='Caption' and 'Current automated validation results' in p.text:p.text=p.text.replace('Current automated validation results','Earlier automated validation results')

# Conservative count includes subheadings and table content, excluding captions.
limits=[1000,2500,2000,2500,2500,1000]; names=['Introduction','Literature Review','Design','Implementation','Evaluation','Conclusion']
counts=[0]*6; current=-1; chapters=[]
for e in d.element.body:
    if e.tag==qn('w:p'):
        text=''.join(e.xpath('.//w:t/text()'))
        styles=e.xpath('./w:pPr/w:pStyle/@w:val');style=styles[0] if styles else ''
        if text.startswith('Chapter '):current+=1;chapters.append(e);continue
        if text=='References':current=6
        if style=='Caption':continue
    elif e.tag==qn('w:tbl'):
        text=' '.join(e.xpath('.//w:t/text()'))
    else:continue
    if 0<=current<6:counts[current]+=len(re.findall(r'\S+',text))
assert all(c<=limit for c,limit in zip(counts,limits)), (counts,limits)
assert sum(counts)<=10500, counts
for p,c,limit,name in zip([old[i] for i in [13,34,65,100,177,213]],counts,limits,names):
    number=names.index(name)+1;p.text=f'{number}. {name} ({c}/{limit} words)'
for row,name,c,limit in zip(d.tables[0].rows[1:],names,counts,limits):
    for cell,value in zip(row.cells,[name,str(c),f'{limit:,}']):cell.text=value
if len(d.tables[0].rows)>7:
    for cell,val in zip(d.tables[0].rows[-1].cells,['Total',str(sum(counts)),'10,500']):cell.text=val
old[10].text=f'Chapter text totals {sum(counts):,} words. Counts include subheadings and table content; exclude chapter titles, captions, title page, abstract and references. The total remains below 10,500 even if the abstract is included.'

# Uniform academic typography and table pagination.
for p in d.paragraphs:
    if p.style.name=='Normal':p.paragraph_format.line_spacing=1.12;p.paragraph_format.space_after=Pt(7)
    if p.style.name=='Heading 1':p.paragraph_format.page_break_before=True
    if p.style.name=='Caption':p.paragraph_format.keep_with_next=False;p.paragraph_format.space_after=Pt(12)
for t in d.tables:
    t.autofit=False
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in t.rows:
        cant=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(cant)
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(5)
                for r in p.runs:r.font.size=Pt(9)
for p in d.paragraphs:
    if p._p.xpath('.//w:drawing'):p.paragraph_format.keep_with_next=True
    if p.text=='The objectives are:':p.paragraph_format.keep_with_next=True
toc=before(old[13],'Contents','Heading 1')
toc.paragraph_format.page_break_before=True
for name,page in zip(names+['References'],[4,6,11,17,49,56,57]):
    p=before(old[13],f'{name}\t{page}')
    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.2))
for section in d.sections:
    footer=section.footer.paragraphs[0]
    footer.alignment=2
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE')
    footer._p.append(field)
for shape in d.inline_shapes:
    if shape.width>Inches(6.5):
        ratio=Inches(6.5)/shape.width;shape.height=int(shape.height*ratio);shape.width=Inches(6.5)
    if shape.height>Inches(7.1):
        ratio=Inches(7.1)/shape.height;shape.width=int(shape.width*ratio);shape.height=Inches(7.1)
d.core_properties.title='KopiBridge AI Final Project Report'
d.core_properties.subject='CM3020 Project Idea 1 Section 4.1'
d.save(OUT)
(REPORTS/'evaluation/report-word-counts.json').write_text(json.dumps({'chapters':dict(zip(names,counts)),'total':sum(counts),'maximum':10500,'figures':len(mapping)},indent=2))
print(json.dumps({'chapters':counts,'total':sum(counts),'figures':len(mapping),'output':str(OUT)}))
